import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

project = Path(__file__).resolve().parents[2]
entries = json.loads((project / "Docs/GronnyDialogues.json").read_text(encoding="utf-8-sig"))
cues = ["Primera bienvenida al acercarse a la mesa. Presentación general para las distintas plantas que se aprenderán.","Inicio de la lección de tomate, después de la bienvenida general.","Antes de plantar la semilla de tomate.","Después de plantar. Explicación inicial del uso de la regadera.","Enseñanza inicial del objetivo y el límite de riego, después de la explicación de la regadera.","Primer riego dentro del rango. Usar durante la enseñanza inicial.","Después de detener el primer riego, cuando ya se puede avanzar.","Al aparecer el brote. Recordatorio breve del nuevo objetivo de agua.","Al aparecer la planta joven. El jugador interpreta el objetivo de riego.","Al aparecer las flores. Gronny deja que el jugador decida la siguiente acción.","Al aparecer los frutos verdes. Último riego con autonomía.","Al madurar el fruto. Primera explicación de cómo cosechar.","Al depositar el tomate maduro en la bandeja de cosecha.","Si la semilla se suelta fuera de la tierra de una maceta válida.","Si se intenta avanzar sin alcanzar el objetivo de agua.","Si se supera el límite de agua de la fase.","Cuando termina el drenaje del exceso de agua.","Si el jugador necesita recordar cómo avanzar con el calendario.","Si la regadera está vacía."]
titles = ["Bienvenida general","Primera lección de tomate","Plantar la semilla","Primer riego","Objetivo de agua","Detener el primer riego","Avanzar por primera vez","Brote","Planta joven","Floración","Frutos verdes","Cosecha","Ciclo completado","Ayuda para plantar","Agua insuficiente","Exceso de agua","Drenaje completado","Ayuda con el calendario","Recargar la regadera"]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.7)
section.left_margin, section.right_margin = Cm(2), Cm(2)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name, size in [("Title", 25), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.space_before = Pt(8 if name.startswith("Heading") else 0)
    if name.startswith("Heading"):
        style.paragraph_format.keep_with_next = True
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn("w:pBdr"))):
            style.element.pPr.remove(border)
doc.core_properties.title = "Guion de voz de Gronny"
doc.core_properties.subject = "19 diálogos en inglés para Cozy VR Gardening"
doc.core_properties.author = ""
doc.core_properties.keywords = "Gronny, KittenTTS, tomate, diálogos"
doc.add_paragraph("Guion de voz de Gronny", "Title")
doc.add_paragraph("Cozy VR Gardening", "Subtitle")
doc.add_paragraph("19 diálogos en inglés para generar la voz de Gronny con KittenTTS. La bienvenida es general y la primera lección enseña a cultivar tomate. Las instrucciones disminuyen durante el crecimiento.")
doc.add_paragraph("Genera un audio por diálogo. Copia solo el texto hablado y conserva el nombre de archivo. Usa una voz cálida, suave y ligeramente juguetona, con la misma configuración en todas las tomas.")
groups = {0: "Bienvenida y enseñanza inicial", 7: "Crecimiento y cosecha", 13: "Ayudas y avisos"}
for i, entry in enumerate(entries):
    if i in groups:
        if i:
            doc.add_page_break()
        doc.add_paragraph(groups[i], "Heading 1")
    doc.add_paragraph(f"{i + 1:02d} {titles[i]}", "Heading 2")
    meta = doc.add_paragraph()
    meta.paragraph_format.keep_with_next = True
    meta.paragraph_format.space_after = Pt(3)
    run = meta.add_run("Archivo " + entry["id"] + ".wav")
    run.bold = True
    run.font.size = Pt(9)
    cue = doc.add_paragraph(cues[i])
    cue.paragraph_format.keep_with_next = True
    cue.paragraph_format.space_after = Pt(4)
    for run in cue.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("4A4A4A")
    speech = doc.add_paragraph(entry["text"])
    speech.paragraph_format.keep_together = True
    speech.paragraph_format.space_after = Pt(7)
doc.add_paragraph("Uso de la ayuda progresiva", "Heading 1")
doc.add_paragraph("Los primeros riegos enseñan el procedimiento. En el último, deja que el jugador decida cuándo detenerse y avanzar. Registra cualquier instrucción de ayuda como asistencia. La cosecha se explica cuando aparece por primera vez.")
doc.add_paragraph("El aviso de exceso puede mantenerse como feedback visual. Si se reproduce una instrucción hablada para resolverlo durante la evaluación, cuenta como ayuda.")
output = project / "Docs/GuionVozGronny.docx"
doc.save(output)
check = Document(output)
texts = [p.text for p in check.paragraphs]
missing = [entry["id"] for entry in entries if entry["text"] not in texts]
if missing:
    raise ValueError("Missing spoken text: " + ", ".join(missing))
print(f"Created document with {len(entries)} complete dialogue entries: {output}")
